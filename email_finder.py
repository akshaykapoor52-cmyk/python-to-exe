import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def find_all_possible_emails():
    # Enter the Excel file path
    file_path = input("Enter the Excel file path: ").strip().strip('"')

    if not Path(file_path).is_file():
        print("Error: Excel file not found.")
        return

    # Enter the person's name (default: Kriti Kapoor)
    person_name = input(
        "Enter the person's full name [Kriti Kapoor]: "
    ).strip()

    if not person_name:
        person_name = "Kriti Kapoor"

    # Normalize the entered name
    parts = person_name.lower().split()

    if len(parts) < 2:
        print("Enter first name and last name.")
        return

    first_name = parts[0]
    last_name = parts[-1]

    # Open the workbook and select the active worksheet
    wb = load_workbook(file_path)
    ws = wb.active

    # Get or create the Search Results worksheet
    if "Search Results" in wb.sheetnames:
        result_ws = wb["Search Results"]

        # Clear existing contents and formatting
        for row in result_ws.iter_rows():
            for cell in row:
                cell.value = None
                cell._style = None
                cell.number_format = "General"
                cell.hyperlink = None
                cell.comment = None
    else:
        result_ws = wb.create_sheet("Search Results")

    # Write the output headers
    headers = ["Search Name", "Email Address", "Match Type"]
    for col, header in enumerate(headers, start=1):
        result_ws.cell(row=1, column=col, value=header)

    out_row = 2
    seen = set()

    # Equivalent to the VBA regular expression
    email_pattern = re.compile(
        r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}",
        re.IGNORECASE
    )

    # Search column A of the active worksheet
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=1).value
        email_text = "" if value is None else str(value)

        for match in email_pattern.finditer(email_text):
            email_addr = match.group(0).lower()
            username = email_addr.split("@", 1)[0]

            # Standardize common username separators
            temp = username
            for separator in [".", "_", "-", "+"]:
                temp = temp.replace(separator, " ")

            temp = temp.replace("'", "")
            tokens = temp.split()
            normalized = "".join(tokens)

            match_type = ""

            # Exact full name without separators
            if normalized == first_name + last_name:
                match_type = "Full name"

            # Exact reversed name without separators
            elif normalized == last_name + first_name:
                match_type = "Reversed name"

            # First name + last initial
            elif normalized == first_name + last_name[:1]:
                match_type = "First name + last initial"

            # First initial + last name
            elif normalized == first_name[:1] + last_name:
                match_type = "First initial + last name"

            else:
                # Match complete adjacent name components
                for i in range(len(tokens) - 1):
                    if (
                        tokens[i] == first_name
                        and tokens[i + 1] == last_name
                    ):
                        match_type = "Full name"
                        break
                    elif (
                        tokens[i] == last_name
                        and tokens[i + 1] == first_name
                    ):
                        match_type = "Reversed name"
                        break

            # Save strong matches only, without duplicate email addresses
            if match_type and email_addr not in seen:
                seen.add(email_addr)

                result_ws.cell(
                    row=out_row, column=1, value=person_name
                )
                result_ws.cell(
                    row=out_row, column=2, value=email_addr
                )
                result_ws.cell(
                    row=out_row, column=3, value=match_type
                )
                out_row += 1

    # Format the header row
    for cell in result_ws[1]:
        cell.font = Font(bold=True)

    # Approximate Excel's AutoFit column widths
    for col in range(1, 4):
        max_length = max(
            (
                len(str(result_ws.cell(row=r, column=col).value or ""))
                for r in range(1, out_row)
            ),
            default=0
        )
        result_ws.column_dimensions[
            get_column_letter(col)
        ].width = max_length + 2

    # Apply filters when results exist
    if out_row > 2:
        result_ws.auto_filter.ref = f"A1:C{out_row - 1}"
        print(f"{out_row - 2} possible matching email addresses found.")
    else:
        result_ws.auto_filter.ref = ""
        print("No strong matching email addresses found.")

    # Save changes back to the original workbook
    wb.save(file_path)
    wb.close()

    print(f"Results saved to: {file_path}")
    print("Worksheet: Search Results")


if __name__ == "__main__":
    find_all_possible_emails()
