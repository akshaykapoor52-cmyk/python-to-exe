import os
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

def run_consolidated_report():
    # Ask for the folder path
    folder_path = input("Enter the folder path containing Excel and CSV files: ").strip()
    
    if not folder_path:
        print("No folder path provided. Exiting.")
        return

    if not folder_path.endswith(('\\', '/')):
        folder_path += os.sep

    if not os.path.isdir(folder_path):
        print(f"Error: The directory '{folder_path}' does not exist.")
        return

    dict_count = {}
    dict_bytes = {}
    file_total = 0
    file_done = 0
    processed_sheets = 0

    # Get all matching files
    files = []
    for filename in os.listdir(folder_path):
        if filename.lower() in ["consolidated report.xlsx", "consolidated report.xlsm"]:
            continue
        if filename.lower().endswith(('.xls', '.xlsx', '.xlsm', '.csv')):
            files.append(filename)
            file_total += 1

    if file_total == 0:
        print(f"No Excel or CSV files were found in: {folder_path}")
        return

    for filename in files:
        file_done += 1
        file_path = os.path.join(folder_path, filename)
        base_filename = os.path.splitext(filename)[0]
        is_custom_file = "_custom" in base_filename.lower()

        try:
            if filename.lower().endswith('.csv'):
                sheets = {base_filename: pd.read_csv(file_path, header=None if is_custom_file else 0)}
            else:
                if is_custom_file:
                    sheets = pd.read_excel(file_path, sheet_name=None, header=None)
                else:
                    sheets = pd.read_excel(file_path, sheet_name=None, header=0)
        except Exception as e:
            print(f"Error reading file {filename}: {e}")
            continue

        for ws_name, df in sheets.items():
            if is_custom_file:
                # Custom files: Row 4 onwards (index 3 in Python), Col A=0, Col B=1, Col C=2
                if df.shape[1] < 3:
                    continue
                
                for r in range(3, len(df)):
                    error_name = str(df.iloc[r, 0]).strip()
                    if error_name and error_name.lower() != 'nan':
                        row_count_val = 0.0
                        row_bytes_val = 0.0
                        
                        c_val = df.iloc[r, 1]
                        s_val = df.iloc[r, 2]
                        
                        if pd.notna(c_val):
                            try:
                                row_count_val = float(c_val)
                            except ValueError:
                                pass
                        if pd.notna(s_val):
                            try:
                                row_bytes_val = float(s_val)
                            except ValueError:
                                pass

                        dict_count[error_name] = dict_count.get(error_name, 0.0) + row_count_val
                        dict_bytes[error_name] = dict_bytes.get(error_name, 0.0) + row_bytes_val
                        processed_sheets += 1
            else:
                # Standard files: Find headers "Count" and "Size"
                count_col = None
                size_col = None
                for col in df.columns:
                    clean_col = str(col).strip().lower()
                    if clean_col == 'count':
                        count_col = col
                    elif clean_col == 'size':
                        size_col = col

                if count_col is not None and size_col is not None:
                    row_count_val = pd.to_numeric(df[count_col], errors='coerce').fillna(0).sum()
                    row_bytes_val = pd.to_numeric(df[size_col], errors='coerce').fillna(0).sum()

                    ws_name_str = str(ws_name)
                    if "ocr" in ws_name_str.lower() or "text" in ws_name_str.lower():
                        error_name = "OCR Exceptions"
                    else:
                        error_name = ws_name_str

                    if "protected" not in error_name.lower() and "encrypted" not in error_name.lower():
                        dict_count[error_name] = dict_count.get(error_name, 0.0) + row_count_val
                        dict_bytes[error_name] = dict_bytes.get(error_name, 0.0) + row_bytes_val
                        processed_sheets += 1

    # Create Report Workbook using openpyxl for formatting and formulas
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Consolidated Report"

    # Headers
    ws_out.cell(row=1, column=1, value="Error Name")
    ws_out.cell(row=1, column=2, value="Count")
    ws_out.cell(row=1, column=3, value="Size (MB)")

    # Header styling
    header_font = openpyxl.styles.Font(bold=True)
    header_fill = openpyxl.styles.PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    for col in range(1, 4):
        cell = ws_out.cell(row=1, column=col)
        cell.font = header_font
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
        cell.fill = header_fill

    output_row = 2
    for key, count_val in dict_count.items():
        ws_out.cell(row=output_ico := output_row, column=1, value=key)
        ws_out.cell(row=output_row, column=2, value=count_val)
        ws_out.cell(row=output_row, column=3, value=dict_bytes[key] / (1024 ** 2))
        output_row += 1

    # Auto-fit columns & number formats
    thin_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin'),
        right=openpyxl.styles.Side(style='thin'),
        top=openpyxl.styles.Side(style='thin'),
        bottom=openpyxl.styles.Side(style='thin')
    )

    for col in range(1, 4):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, output_row):
            val = ws_out.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
                if row > 1:
                    cell = ws_out.cell(row=row, column=col)
                    cell.border = thin_border
                    if col == 2:
                        cell.number_format = '#,##0'
                    elif col == 3:
                        cell.number_format = '#,##0.00'
        ws_out.column_dimensions[col_letter].width = max(max_len + 3, 12)

    if output_row > 2:
        for col in range(1, 4):
            ws_out.cell(row=1, column=col).border = thin_border

    # Grand Total row
    ws_out.cell(row=output_row, column=1, value="Grand Total").font = openpyxl.styles.Font(bold=True)
    ws_out.cell(row=output_row, column=2, value=f"=SUM(B2:B{output_row - 1})").font = openpyxl.styles.Font(bold=True)
    ws_out.cell(row=output_row, column=3, value=f"=SUM(C2:C{output_row - 1})").font = openpyxl.styles.Font(bold=True)

    for col in range(1, 4):
        cell = ws_out.cell(row=output_row, column=col)
        cell.border = thin_border

    output_file_path = os.path.join(folder_path, "Consolidated report.xlsx")
    wb_out.save(output_file_path)

    print("\nReport Complete!")
    print(f"Files processed: {file_done}")
    print(f"Sheets processed: {processed_sheets}")
    print(f"Location: {output_file_path}")

if __name__ == "__main__":
    run_consolidated_report()
