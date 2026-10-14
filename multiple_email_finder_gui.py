import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def sanitize_sheet_name(name):
    """Sanitize string to comply with Excel sheet name rules (max 31 chars, no invalid characters)."""
    invalid_chars = ['\\', '/', '?', '*', ':', '[', ']']
    for char in invalid_chars:
        name = name.replace(char, '')
    return name[:31].strip() or "Sheet"


class EmailFinderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Email Finder & Excel Multi-Sheet Generator")
        self.root.geometry("620x520")
        self.root.minsize(580, 480)

        # Style configuration
        self.style = ttk.Style()
        self.style.theme_use("clam")

        # Main Container Frame
        main_frame = ttk.Frame(root, padding=16)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # Title Label
        title_label = ttk.Label(
            main_frame, 
            text="Excel Email Extractor & Matcher", 
            font=("Segoe UI", 16, "bold")
        )
        title_label.pack(anchor=tk.W, pady=(0, 15))

        # --- Directory Selection Frame ---
        dir_frame = ttk.LabelFrame(main_frame, text=" Target Directory ", padding=12)
        dir_frame.pack(fill=tk.X, pady=(0, 12))

        self.dir_var = tk.StringVar(value=str(Path.cwd()))
        dir_entry = ttk.Entry(dir_frame, textvariable=self.dir_var, font=("Segoe UI", 10))
        dir_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))

        browse_btn = ttk.Button(dir_frame, text="Browse...", command=self.browse_directory)
        browse_btn.pack(side=tk.RIGHT)

        # --- Names Input Frame ---
        names_frame = ttk.LabelFrame(main_frame, text=" People Names (Comma Separated) ", padding=12)
        names_frame.pack(fill=tk.X, pady=(0, 12))

        self.names_var = tk.StringVar(value="Kriti Kapoor")
        names_entry = ttk.Entry(names_frame, textvariable=self.names_var, font=("Segoe UI", 10))
        names_entry.pack(fill=tk.X, expand=True)
        
        hint_label = ttk.Label(
            names_frame, 
            text="Example: Kriti Kapoor, John Doe, Sarah Smith", 
            font=("Segoe UI", 8, "italic"),
            foreground="#666666"
        )
        hint_label.pack(anchor=tk.W, pady=(4, 0))

        # --- Action & Progress Frame ---
        action_frame = ttk.Frame(main_frame)
        action_frame.pack(fill=tk.X, pady=(0, 12))

        self.run_btn = ttk.Button(action_frame, text="Run Search & Generate Excel", command=self.start_processing)
        self.run_btn.pack(side=tk.LEFT, ipadx=10, ipady=5)

        self.progress_bar = ttk.Progressbar(action_frame, mode='indeterminate', length=250)
        # Progress bar will be packed when running

        # --- Log / Console Frame ---
        log_frame = ttk.LabelFrame(main_frame, text=" Execution Output Log ", padding=10)
        log_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 5))

        self.log_text = tk.Text(log_frame, height=10, wrap=tk.WORD, font=("Consolas", 9), background="#f9f9f9")
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.configure(yscrollcommand=scrollbar.set)

    def browse_directory(self):
        chosen_dir = filedialog.askdirectory(initialdir=self.dir_var.get())
        if chosen_dir:
            self.dir_var.set(chosen_dir)

    def log(self, message):
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)

    def start_processing(self):
        target_dir_str = self.dir_var.get().strip()
        names_input = self.names_var.get().strip()

        if not target_dir_str or not names_input:
            messagebox.showerror("Error", "Please provide both a valid target directory and names.")
            return

        self.run_btn.state(["disabled"])
        self.progress_bar.pack(side=tk.RIGHT, padx=5)
        self.progress_bar.start(10)
        self.log_text.delete("1.0", tk.END)

        # Run processing in a background thread to prevent UI freezing
        threading.Thread(target=self.process_files, args=(target_dir_str, names_input), daemon=True).start()

    def process_files(self, target_dir_str, names_input):
        try:
            target_dir = Path(target_dir_str)
            if not target_dir.is_dir():
                self.root.after(0, lambda: messagebox.showerror("Error", f"Directory not found -> {target_dir}"))
                return

            excel_files = list(target_dir.glob("*.xlsx"))
            excel_files = [f for f in excel_files if f.name != "Master_Email_Search_Results.xlsx"]

            if not excel_files:
                self.root.after(0, lambda: messagebox.showwarning("Warning", f"No Excel (.xlsx) files found in: {target_dir}"))
                return

            self.log(f"Found {len(excel_files)} Excel file(s) to process.")

            raw_names = [n.strip() for n in names_input.split(",") if n.strip()]
            processed_persons = []

            for full_name in raw_names:
                parts = full_name.lower().split()
                if len(parts) >= 2:
                    processed_persons.append({
                        "original_name": full_name,
                        "first_name": parts[0],
                        "last_name": parts[-1]
                    })

            if not processed_persons:
                self.root.after(0, lambda: messagebox.showerror("Error", "Please include valid full names (first and last name)."))
                return

            email_pattern = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)
            all_results = []
            seen = set()

            for file_path in excel_files:
                self.log(f"Processing file: {file_path.name}...")
                try:
                    wb = load_workbook(file_path, read_only=True, data_only=True)
                    ws = wb.active
                except Exception as e:
                    self.log(f"  -> Error reading {file_path.name}: {e}")
                    continue

                file_emails = []
                for row in range(1, ws.max_row + 1):
                    val = ws.cell(row=row, column=1).value
                    if val:
                        file_emails.append(str(val))
                wb.close()

                for email_text in file_emails:
                    for match in email_pattern.finditer(email_text):
                        email_addr = match.group(0).lower()
                        username = email_addr.split("@", 1)[0]

                        temp = username
                        for separator in [".", "_", "-", "+"]:
                            temp = temp.replace(separator, " ")
                        temp = temp.replace("'", "")
                        tokens = temp.split()
                        normalized = "".join(tokens)

                        for person in processed_persons:
                            first_name = person["first_name"]
                            last_name = person["last_name"]
                            person_name = person["original_name"]
                            match_type = ""

                            if normalized == first_name + last_name:
                                match_type = "Full name"
                            elif normalized == last_name + first_name:
                                match_type = "Reversed name"
                            elif normalized == first_name + last_name[:1]:
                                match_type = "First name + last initial"
                            elif normalized == first_name[:1] + last_name:
                                match_type = "First initial + last name"
                            else:
                                for i in range(len(tokens) - 1):
                                    if tokens[i] == first_name and tokens[i + 1] == last_name:
                                        match_type = "Full name"
                                        break
                                    elif tokens[i] == last_name and tokens[i + 1] == first_name:
                                        match_type = "Reversed name"
                                        break

                            if match_type:
                                unique_key = (file_path.name, email_addr, person_name)
                                if unique_key not in seen:
                                    seen.add(unique_key)
                                    all_results.append({
                                        "file_name": file_path.name,
                                        "search_name": person_name,
                                        "email": email_addr,
                                        "match_type": match_type
                                    })

            # Create Workbook with Multi-Sheet Structure
            output_wb = Workbook()
            summary_ws = output_wb.active
            summary_ws.title = "All Results"

            headers = ["File Name", "Search Name", "Email Address", "Match Type"]
            for col, header in enumerate(headers, start=1):
                summary_ws.cell(row=1, column=col, value=header)

            sum_row = 2
            for res in all_results:
                summary_ws.cell(row=sum_row, column=1, value=res["file_name"])
                summary_ws.cell(row=sum_row, column=2, value=res["search_name"])
                summary_ws.cell(row=sum_row, column=3, value=res["email"])
                summary_ws.cell(row=sum_row, column=4, value=res["match_type"])
                sum_row += 1

            for cell in summary_ws[1]:
                cell.font = Font(bold=True)

            for col in range(1, 5):
                max_len = max((len(str(summary_ws.cell(row=r, column=col).value or "")) for r in range(1, sum_row)), default=0)
                summary_ws.column_dimensions[get_column_letter(col)].width = max_len + 3

            if sum_row > 2:
                summary_ws.auto_filter.ref = f"A1:D{sum_row - 1}"

            for person in processed_persons:
                person_name = person["original_name"]
                sheet_title = sanitize_sheet_name(person_name)
                person_ws = output_wb.create_sheet(title=sheet_title)

                p_headers = ["File Name", "Email Address", "Match Type"]
                for col, header in enumerate(p_headers, start=1):
                    person_ws.cell(row=1, column=col, value=header)

                person_results = [r for r in all_results if r["search_name"] == person_name]
                p_row = 2
                for res in person_results:
                    person_ws.cell(row=p_row, column=1, value=res["file_name"])
                    person_ws.cell(row=p_row, column=2, value=res["email"])
                    person_ws.cell(row=p_row, column=3, value=res["match_type"])
                    p_row += 1

                for cell in person_ws[1]:
                    cell.font = Font(bold=True)

                for col in range(1, 4):
                    max_len = max((len(str(person_ws.cell(row=r, column=col).value or "")) for r in range(1, p_row)), default=0)
                    person_ws.column_dimensions[get_column_letter(col)].width = max_len + 3

                if p_row > 2:
                    person_ws.auto_filter.ref = f"A1:C{p_row - 1}"

            output_file_path = target_dir / "Master_Email_Search_Results.xlsx"
            output_wb.save(output_file_path)
            output_wb.close()

            found_count = sum_row - 2
            self.log(f"\nDone! Found {found_count} matching email record(s).")
            self.log(f"Results saved to: {output_file_path}")
            self.root.after(0, lambda: messagebox.showinfo("Success", f"Processing complete! Found {found_count} matches.\nSaved to Master_Email_Search_Results.xlsx"))

        except Exception as ex:
            self.log(f"Critical Error: {ex}")
            self.root.after(0, lambda: messagebox.showerror("Error", str(ex)))
        finally:
            self.root.after(0, self.reset_ui)

    def reset_ui(self):
        self.progress_bar.stop()
        self.progress_bar.pack_forget()
        self.run_btn.state(["!disabled"])


if __name__ == "__main__":
    root = tk.Tk()
    app = EmailFinderApp(root)
    root.mainloop()
