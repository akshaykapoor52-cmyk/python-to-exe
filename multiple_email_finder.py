import re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def sanitize_sheet_name(name):
    """Sanitize string to comply with Excel sheet name rules (max 31 chars, no invalid characters)."""
    invalid_chars = ['\\', '/', '?', '*', ':', '[', ']']
    for char in invalid_chars:
        name = name.replace(char, '')
    # Truncate to 31 characters max
    return name[:31].strip() or "Sheet"


def find_all_possible_emails():
    # 1. Get directory path containing Excel files
    dir_input = input("Enter the directory path containing Excel files (leave blank for current directory): ").strip().strip('"')
    target_dir = Path(dir_input) if dir_input else Path.cwd()

    if not target_dir.is_dir():
        print(f"Error: Directory not found -> {target_dir}")
        return

    # Find all Excel files in the directory
    excel_files = list(target_dir.glob("*.xlsx"))
    # Exclude any previously generated master result file to avoid recursive reading
    excel_files = [f for f in excel_files if f.name != "Master_Email_Search_Results.xlsx"]

    if not excel_files:
        print(f"No Excel (.xlsx) files found in: {target_dir}")
        return

    print(f"Found {len(excel_files)} Excel file(s) to process.")

    # 2. Get multiple names separated by commas
    names_input = input("Enter full names separated by commas [Kriti Kapoor]: ").strip()
    if not names_input:
        names_input = "Kriti Kapoor"

    raw_names = [n.strip() for n in names_input.split(",") if n.strip()]
    processed_persons = []

    for full_name in raw_names:
        parts = full_name.lower().split()
        if len(parts) >= 2:
            processed_persons.append({
                "original_name": full_name,
                "first_name": parts[0],
                "last_name": parts[-1]
            })

    if not processed_persons:
        print("Error: No valid names provided. Please include both first and last names.")
        return

    # Email regex pattern
    email_pattern = re.compile(
        r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}",
        re.IGNORECASE
    )

    all_results = []
    seen = set()  # Track unique combinations of (file_name, email_addr, person_name)

    # 3. Process each Excel file
    for file_path in excel_files:
        print(f"Processing: {file_path.name}...")
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
        except Exception as e:
            print(f"  -> Could not read file {file_path.name}. Error: {e}")
            continue

        # Extract rows from column A
        file_emails = []
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row=row, column=1).value
            if value:
                file_emails.append(str(value))
        wb.close()

        # Check each email against each person
        for email_text in file_emails:
            for match in email_pattern.finditer(email_text):
                email_addr = match.group(0).lower()
                username = email_addr.split("@", 1)[0]

                # Standardize common username separators
                temp = username
                for separator in [".", "_", "-", "+"]:
                    temp = temp.replace(separator, " ")
                temp = temp.replace("'", "")
                tokens = temp.split()
                normalized = "".join(tokens)

                for person in processed_persons:
                    first_name = person["first_name"]
                    last_name = person["last_name"]
                    person_name = person["original_name"]
                    match_type = ""

                    # Exact full name without separators
                    if normalized == first_name + last_name:
                        match_type = "Full name"
                    # Exact reversed name without separators
                    elif normalized == last_name + first_name:
                        match_type = "Reversed name"
                    # First name + last initial
                    elif normalized == first_name + last_name[:1]:
                        match_type = "First name + last initial"
                    # First initial + last name
                    elif normalized == first_name[:1] + last_name:
                        match_type = "First initial + last name"
                    else:
                        # Match complete adjacent name components
                        for i in range(len(tokens) - 1):
                            if tokens[i] == first_name and tokens[i + 1] == last_name:
                                match_type = "Full name"
                                break
                            elif tokens[i] == last_name and tokens[i + 1] == first_name:
                                match_type = "Reversed name"
                                break

                    if match_type:
                        unique_key = (file_path.name, email_addr, person_name)
                        if unique_key not in seen:
                            seen.add(unique_key)
                            all_results.append({
                                "file_name": file_path.name,
                                "search_name": person_name,
                                "email": email_addr,
                                "match_type": match_type
                            })

    # 4. Create Master Results Workbook with Multi-Sheet Structure
    output_wb = Workbook()
    
    # Setup default sheet as "All Results" Summary
    summary_ws = output_wb.active
    summary_ws.title = "All Results"

    headers = ["File Name", "Search Name", "Email Address", "Match Type"]
    for col, header in enumerate(headers, start=1):
        summary_ws.cell(row=1, column=col, value=header)

    # Populate summary sheet
    sum_row = 2
    for res in all_results:
        summary_ws.cell(row=sum_row, column=1, value=res["file_name"])
        summary_ws.cell(row=sum_row, column=2, value=res["search_name"])
        summary_ws.cell(row=sum_row, column=3, value=res["email"])
        summary_ws.cell(row=sum_row, column=4, value=res["match_type"])
        sum_row += 1

    # Format summary sheet header
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)

    # Auto-fit columns for summary sheet
    for col in range(1, 5):
        max_length = max(
            (len(str(summary_ws.cell(row=r, column=col).value or "")) for r in range(1, sum_row)),
            default=0
        )
        summary_ws.column_dimensions[get_column_letter(col)].width = max_length + 3

    if sum_row > 2:
        summary_ws.auto_filter.ref = f"A1:D{sum_row - 1}"

    # 5. Create Individual Sheets for Each Person
    for person in processed_persons:
        person_name = person["original_name"]
        sheet_title = sanitize_sheet_name(person_name)
        
        # Create a new sheet for this person
        person_ws = output_wb.create_sheet(title=sheet_title)
        
        # Write headers (excluding 'Search Name' since it's redundant on a dedicated sheet)
        p_headers = ["File Name", "Email Address", "Match Type"]
        for col, header in enumerate(p_headers, start=1):
            person_ws.cell(row=1, column=col, value=header)
            
        # Filter results belonging to this specific person
        person_results = [r for r in all_results if r["search_name"] == person_name]
        
        p_row = 2
        for res in person_results:
            person_ws.cell(row=p_row, column=1, value=res["file_name"])
            person_ws.cell(row=p_row, column=2, value=res["email"])
            person_ws.cell(row=p_row, column=3, value=res["match_type"])
            p_row += 1
            
        # Format person sheet header
        for cell in person_ws[1]:
            cell.font = Font(bold=True)
            
        # Auto-fit columns for person sheet
        for col in range(1, 4):
            max_length = max(
                (len(str(person_ws.cell(row=r, column=col).value or "")) for r in range(1, p_row)),
                default=0
            )
            person_ws.column_dimensions[get_column_letter(col)].width = max_length + 3
            
        if p_row > 2:
            person_ws.auto_filter.ref = f"A1:C{p_row - 1}"

    # Save output file
    output_file_path = target_dir / "Master_Email_Search_Results.xlsx"
    output_wb.save(output_file_path)
    output_wb.close()

    if sum_row > 2:
        print(f"\nSuccess! Found {sum_row - 2} matching email record(s).")
    else:
        print("\nNo matching email addresses found across the scanned files.")

    print(f"Results saved with individual sheets per person to: {output_file_path}")


if __name__ == "__main__":
    find_all_possible_emails()
